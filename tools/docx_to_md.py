"""Convert a .docx file to Markdown.

Usage:
    python tools/docx_to_md.py <input.docx> [output.md]

If output is omitted, the .md file is written next to the input with the same name.
Handles headings, paragraphs, bold/italic runs, hyperlinks, bulleted/numbered lists
and tables, preserving document order. Requires: pip install python-docx
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(doc):
    """Yield paragraphs and tables in document order."""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def escape_md(text):
    return re.sub(r"([\\`*_])", r"\\\1", text)


def format_run_text(text, bold, italic):
    if not text.strip():
        return text
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = escape_md(text.strip())
    if bold and italic:
        core = f"***{core}***"
    elif bold:
        core = f"**{core}**"
    elif italic:
        core = f"*{core}*"
    return f"{lead}{core}{trail}"


def paragraph_inline(par):
    """Render a paragraph's runs (and hyperlinks) to inline Markdown."""
    parts = []
    for child in par._p.iterchildren():
        if child.tag == qn("w:r"):
            parts.append(run_xml_to_md(child))
        elif child.tag == qn("w:hyperlink"):
            text = "".join(run_xml_to_md(r) for r in child.iter(qn("w:r")))
            rid = child.get(qn("r:id"))
            url = None
            if rid and rid in par.part.rels:
                url = par.part.rels[rid].target_ref
            parts.append(f"[{text}]({url})" if url else text)
    return merge_markers("".join(parts)).strip()


def run_xml_to_md(r):
    text = ""
    for node in r.iterchildren():
        if node.tag == qn("w:t"):
            text += node.text or ""
        elif node.tag == qn("w:tab"):
            text += " "
        elif node.tag in (qn("w:br"), qn("w:cr")):
            text += "<br>"
    rpr = r.find(qn("w:rPr"))
    bold = italic = False
    if rpr is not None:
        b = rpr.find(qn("w:b"))
        i = rpr.find(qn("w:i"))
        bold = b is not None and b.get(qn("w:val")) not in ("0", "false")
        italic = i is not None and i.get(qn("w:val")) not in ("0", "false")
    return format_run_text(text, bold, italic)


def merge_markers(text):
    """Join adjacent formatted runs, e.g. '**a****b**' -> '**ab**'."""
    return text.replace("******", "").replace("****", "")


def heading_level(par):
    name = (par.style.name or "").lower() if par.style is not None else ""
    if name == "title":
        return 1
    m = re.match(r"heading\s*(\d)", name)
    if m:
        return min(int(m.group(1)) + 1, 6) if has_title else int(m.group(1))
    return 0


def list_info(par):
    """Return (ordered, level) if the paragraph is a list item, else None."""
    ppr = par._p.pPr
    style = (par.style.name or "").lower() if par.style is not None else ""
    numpr = ppr.numPr if ppr is not None else None
    if numpr is None and "list" not in style:
        return None
    level = 0
    if numpr is not None and numpr.ilvl is not None:
        level = int(numpr.ilvl.val)
    ordered = "number" in style
    if numpr is not None and numpr.numId is not None:
        ordered = is_ordered_numbering(par, int(numpr.numId.val), level) or ordered
    return ordered, level


def is_ordered_numbering(par, num_id, level):
    try:
        numbering = par.part.numbering_part.element
    except (NotImplementedError, KeyError):
        return False
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) != str(num_id):
            continue
        abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        for absn in numbering.findall(qn("w:abstractNum")):
            if absn.get(qn("w:abstractNumId")) != abs_id:
                continue
            for lvl in absn.findall(qn("w:lvl")):
                if lvl.get(qn("w:ilvl")) == str(level):
                    fmt = lvl.find(qn("w:numFmt"))
                    return fmt is not None and fmt.get(qn("w:val")) not in ("bullet", "none")
    return False


def cell_text(cell):
    lines = [paragraph_inline(p) for p in cell.paragraphs]
    return "<br>".join(line for line in lines if line).replace("|", "\\|")


def table_to_md(table):
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            # Merged cells repeat the same underlying element; render once.
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(cell_text(cell))
        rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


has_title = False


def convert(src, dst):
    global has_title
    doc = Document(src)
    has_title = any(
        p.style is not None and (p.style.name or "").lower() == "title" for p in doc.paragraphs
    )
    blocks = []
    prev_list = False
    for block in iter_block_items(doc):
        if isinstance(block, Table):
            blocks.append(table_to_md(block))
            prev_list = False
            continue
        text = paragraph_inline(block)
        if not text:
            prev_list = False
            continue
        level = heading_level(block)
        info = list_info(block)
        if level:
            plain = re.sub(r"^\*+|\*+$", "", text)
            blocks.append("#" * level + " " + plain)
            prev_list = False
        elif info:
            ordered, lvl = info
            item = "  " * lvl + ("1. " if ordered else "- ") + text
            if prev_list:
                blocks[-1] += "\n" + item
            else:
                blocks.append(item)
            prev_list = True
        else:
            blocks.append(text)
            prev_list = False
    Path(dst).write_text("\n\n".join(blocks).strip() + "\n", encoding="utf-8")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".md")
    convert(src, dst)
    print(f"Wrote {dst}")


if __name__ == "__main__":
    main()
